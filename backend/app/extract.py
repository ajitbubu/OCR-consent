"""Page rendering, native extraction and provider-based OCR for the worker."""
import json
import sys
import warnings
from pathlib import Path
from PIL import Image, ImageOps, ImageDraw, ImageFont
from pypdf import PdfReader
import pypdfium2 as pdfium
from .categorize import categorize
from .ocr import recognize

MAX_PAGES = 100
Image.MAX_IMAGE_PIXELS = 25_000_000
warnings.simplefilter('error', Image.DecompressionBombWarning)


def extract(source: Path, output: Path) -> dict:
    output.mkdir(exist_ok=True, parents=True)
    pages, notes = [], []
    def add(image, native='', widgets=None):
        n = len(pages)+1
        if n > MAX_PAGES: raise ValueError('Document exceeds the 100-page pilot limit.')
        image = ImageOps.exif_transpose(image).convert('RGB')
        ocr_result = recognize(image)
        preview_path = output / f'{n}.png'
        preview = image.copy(); preview.thumbnail((1300, 1800)); preview.save(preview_path)
        method = f'Native + {ocr_result.method}' if native.strip() else ocr_result.method
        pages.append(dict(number=n, text=ocr_result.text, native_text=native,
                          method=method, confidence=ocr_result.confidence,
                          widgets=widgets or [], image_path=str(preview_path.resolve())))
    suffix = source.suffix.lower()
    if suffix == '.pdf':
        reader = PdfReader(source)
        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception as exc:
                raise ValueError("Password-protected PDFs require an empty password in this pilot.") from exc
            if not unlocked:
                raise ValueError("Password-protected PDFs require an empty password in this pilot.")
        if len(reader.pages) > MAX_PAGES: raise ValueError('Document exceeds the 100-page pilot limit.')
        doc = pdfium.PdfDocument(str(source)); doc.init_forms()
        for i, page in enumerate(reader.pages):
            widgets = []
            for ref in page.get('/Annots', []):
                w = ref.get_object()
                if w.get('/Subtype') != '/Widget': continue
                parent = w.get('/Parent'); f = parent.get_object() if parent else w
                value = w.get('/V', f.get('/V'))
                if value not in [None, '', '/Off']:
                    widgets.append(dict(name=str(f.get('/T', 'Unnamed')), value=str(value), appearance=str(w.get('/AS', '')), rect=[float(x) for x in w.get('/Rect', [])]))
            p = doc[i]
            width, height = p.get_size()
            if width * height * 4 > 25_000_000: raise ValueError('PDF page exceeds the render size limit.')
            bitmap = p.render(scale=2, may_draw_forms=True)
            native = page.extract_text() or ''
            try:
                layout = page.extract_text(extraction_mode='layout') or ''
            except Exception:
                layout = ''
            if layout.strip() and layout.strip() != native.strip():
                native = f'{native}\n\n{layout}'
            add(bitmap.to_pil(), native, widgets)
            bitmap.close(); p.close()
        doc.close()
    elif suffix == '.docx':
        import zipfile
        from docx import Document
        import textwrap
        with zipfile.ZipFile(source) as z:
            if sum(i.file_size for i in z.infolist()) > 100_000_000: raise ValueError('Expanded Word document is too large.')
            if any(i.filename.endswith('vbaProject.bin') for i in z.infolist()): raise ValueError('Macro-bearing documents are not accepted.')
        d = Document(source)
        text = '\n'.join(p.text for p in d.paragraphs) + '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows)
        notes.append('Word preview is an extracted-text view, not the original layout. Embedded images are processed separately.')
        lines = []
        for line in text.splitlines(): lines.extend(textwrap.wrap(line, 90) or [''])
        for start in range(0, max(1, len(lines)), 55):
            native = '\n'.join(lines[start:start+55]); image = Image.new('RGB', (1100, 1500), 'white')
            ImageDraw.Draw(image).multiline_text((45,45), native, fill='black', font=ImageFont.load_default(size=19), spacing=5)
            n = len(pages)+1; image.save(output / f'{n}.png')
            pages.append(dict(number=n, text=native, native_text=native, method='Word text',
                              confidence=None, widgets=[], image_path=str((output / f'{n}.png').resolve())))
        with zipfile.ZipFile(source) as z:
            import io
            for name in z.namelist():
                if name.startswith('word/media/') and not name.endswith('/'):
                    try:
                        with Image.open(io.BytesIO(z.read(name))) as im: add(im)
                    except (OSError, ValueError) as exc: notes.append(f'An embedded image could not be processed: {type(exc).__name__}.')
    else:
        with Image.open(source) as im:
            if getattr(im, 'n_frames', 1) > MAX_PAGES: raise ValueError('Image exceeds the 100-frame limit.')
            for i in range(getattr(im, 'n_frames', 1)):
                im.seek(i); add(im.copy())
    if not pages: raise ValueError('No readable pages found.')
    result = categorize(pages)
    if not any(p['text'].strip() or p['native_text'].strip() for p in pages): notes.append('No text detected. Review the original image.')
    if any(p['widgets'] for p in pages): notes.append('PDF contains stored field values. Saved selections alone do not identify who chose them.')
    result['warnings'] += notes
    result['pages'] = pages
    (output / 'result.json').write_text(json.dumps(result), encoding='utf-8')
    return result

if __name__ == '__main__':
    extract(Path(sys.argv[1]), Path(sys.argv[2]))
